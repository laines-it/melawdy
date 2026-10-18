"""Текст → DOCX bytes."""

from io import BytesIO

from docx import Document


def text_to_docx(text: str) -> bytes:
    doc = Document()
    for block in (text or "").split("\n\n"):
        block = block.strip("\n")
        if not block:
            continue
        lines = block.split("\n")
        p = doc.add_paragraph(lines[0])
        for line in lines[1:]:
            p.add_run("\n" + line)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()