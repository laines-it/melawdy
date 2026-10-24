import json
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock

import pytest
from docx import Document
from fastapi import HTTPException

from backend import cabinet, public
from backend.models import DocumentTemplate, DocumentType, User, UserDocument
from backend.render import render_flat_body, render_structure

ROOT = Path(__file__).resolve().parents[3]
VECTORS = json.loads((ROOT / "web/tests/fixtures/document-render-cases.json").read_text(encoding="utf-8"))


@pytest.mark.parametrize("vector", VECTORS, ids=lambda item: item["name"])
def test_browser_preview_and_server_render_share_contract(vector):
    template, params = vector["template"], vector["params"]
    structure = template.get("structure_json", {})
    text = (render_structure(structure, params) if structure.get("blocks")
            else render_flat_body(template.get("body", ""), params))
    assert text == vector["expected"]


def test_public_template_exposes_constructor_metadata_without_internal_file_path():
    template = DocumentTemplate(id=4, slug="coauthors", title="Соавторы", document_type=DocumentType.CONTRACT,
                                body="{{city}}", schema_json={"sections": []}, structure_json={"blocks": []},
                                file_path="private/file.docx")
    session = Mock()
    session.scalar.return_value = template
    result = public.read_template(4, session)
    assert result["id"] == "4"
    assert result["slug"] == "coauthors"
    assert result["schema_json"] == {"sections": []}
    assert result["structure_json"] == {"blocks": []}
    assert "file_path" not in result


def test_download_renders_saved_parameters_without_creating_duplicates():
    session = Mock()
    session.scalar.return_value = UserDocument(id=7, user_id=8, template_id=4,
                                               content_json={"city": "Москва", "proxy": False})
    session.get.return_value = DocumentTemplate(id=4, body="ignored", structure_json={"blocks": [
        {"template": "г. {{city}}"}, {"when": "proxy == true", "template": "Представитель"},
    ]})
    result = cabinet.download_document(7, User(id=8), session)
    paragraphs = Document(BytesIO(result.body)).paragraphs
    assert [p.text for p in paragraphs] == ["г. Москва"]
    assert result.headers["cache-control"] == "no-store"
    assert "document-7.docx" in result.headers["content-disposition"]
    session.add.assert_not_called()
    session.commit.assert_not_called()


def test_download_of_someone_elses_document_is_not_found():
    session = Mock()
    session.scalar.return_value = None
    with pytest.raises(HTTPException) as error:
        cabinet.download_document(7, User(id=9), session)
    assert error.value.status_code == 404
    session.get.assert_not_called()


@pytest.mark.parametrize("template_id,template,expected", [
    (None, None, 404), (4, None, 404),
    (4, DocumentTemplate(body="", structure_json={}), 422),
])
def test_missing_or_empty_download_has_an_explicit_error(template_id, template, expected):
    session = Mock()
    session.scalar.return_value = UserDocument(id=7, user_id=8, template_id=template_id, content_json={})
    session.get.return_value = template
    with pytest.raises(HTTPException) as error:
        cabinet.download_document(7, User(id=8), session)
    assert error.value.status_code == expected
